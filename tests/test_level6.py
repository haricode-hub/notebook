"""Level 6 artifact recipes, repair behavior, ownership, and exports."""

import time
from collections.abc import AsyncIterator
from io import BytesIO
from typing import Any

import httpx
import pytest
from docx import Document
from opennotebook_api.routers import artifacts as artifacts_router
from opennotebook_core.artifacts import (
    RECIPES_REGISTRY,
    ArtifactDocument,
    ArtifactSection,
    FakeArtifactLLM,
    Flashcard,
    FlashcardsArtifact,
    MindMapArtifact,
    MindMapNode,
    QuizArtifact,
    QuizQuestion,
    generate_artifact,
)
from opennotebook_core.config import Settings
from opennotebook_core.models import Chunk, Source
from opennotebook_worker.main import generate_text_artifact
from sqlalchemy import text
from test_level1 import auth
from test_level3_chat import seeded_book
from test_level5 import QueueFake

pytest_plugins = ["test_level1"]


class ScriptedLLM:
    def __init__(self, responses: list[str]) -> None:
        self.responses = responses
        self.calls: list[list[dict[str, str]]] = []

    async def stream(self, messages: list[dict[str, str]]) -> AsyncIterator[str]:
        self.calls.append(messages)
        yield self.responses.pop(0)


def fixture_document(recipe: str, ref: int = 1) -> str:
    if recipe == "mindmap":
        return MindMapArtifact(
            title="Mind map",
            root=MindMapNode(
                label="Launch",
                summary=f"The launch code is BLUEBIRD-42 [S{ref}].",
                source_refs=[ref],
            ),
        ).model_dump_json()
    if recipe == "flashcards":
        return FlashcardsArtifact(
            title="Flashcards",
            cards=[
                Flashcard(
                    front="What is the launch code?", back=f"BLUEBIRD-42 [S{ref}]", source_ref=ref
                )
            ],
        ).model_dump_json()
    if recipe == "quiz":
        return QuizArtifact(
            title="Quiz",
            questions=[
                QuizQuestion(
                    question="What is the launch code?",
                    options=["BLUEBIRD-42", "RED-99", "GREEN-7", "None"],
                    answer_index=0,
                    explanation=f"The source says BLUEBIRD-42. [S{ref}]",
                    source_ref=ref,
                )
                for _ in range(10)
            ],
        ).model_dump_json()
    return ArtifactDocument(
        recipe=recipe,
        title=f"{RECIPES_REGISTRY[recipe].title} for launch",
        sections=[
            ArtifactSection(
                heading="Launch finding",
                markdown=f"The launch code is BLUEBIRD-42 [S{ref}].",
                citation_refs=[ref],
            )
        ],
    ).model_dump_json()


def sample_chunk(
    number: int = 1, *, source_id: str = "evidence", notebook_id: str = "book"
) -> Chunk:
    text = "The launch code is BLUEBIRD-42. It identifies the research mission."
    return Chunk(
        id=f"chunk-{number}",
        source_id=source_id,
        notebook_id=notebook_id,
        ordinal=number,
        text=text,
        page_start=number,
        char_start=0,
        char_end=len(text),
        heading_path=[],
        embedding=[0.0] * 768,
        metadata_json={},
    )


@pytest.mark.asyncio
@pytest.mark.parametrize("recipe", list(RECIPES_REGISTRY))
async def test_each_recipe_accepts_grounded_fixture(recipe: str) -> None:
    spy = ScriptedLLM([fixture_document(recipe)])
    document, citations = await generate_artifact(
        recipe, [sample_chunk()], {"evidence": "Launch.txt"}, {}, spy
    )
    assert document.recipe == recipe
    assert citations[0]["chunk_id"] == "chunk-1"
    assert citations[0]["page_start"] == 1
    assert len(spy.calls) == 1


@pytest.mark.asyncio
async def test_invalid_json_gets_exactly_one_repair_and_fabricated_refs_fail() -> None:
    chunk = sample_chunk()
    repaired = ScriptedLLM(["not json", fixture_document("study_guide")])
    document, citations = await generate_artifact(
        "study_guide", [chunk], {"evidence": "Launch.txt"}, {}, repaired
    )
    assert isinstance(document, ArtifactDocument)
    assert document.sections[0].citation_refs == [1]
    assert len(citations) == 1 and len(repaired.calls) == 2
    assert repaired.calls[1][-1]["role"] == "user"

    fabricated = ScriptedLLM([fixture_document("study_guide", 99)] * 2)
    with pytest.raises(ValueError, match="after repair"):
        await generate_artifact("study_guide", [chunk], {"evidence": "Launch.txt"}, {}, fabricated)
    assert len(fabricated.calls) == 2


@pytest.mark.asyncio
async def test_map_reduce_fifty_pages_all_recipes_under_ninety_seconds() -> None:
    chunks = [sample_chunk(number) for number in range(1, 51)]
    for recipe in RECIPES_REGISTRY:
        start = time.perf_counter()
        document, citations = await generate_artifact(
            recipe,
            chunks,
            {"evidence": "50 pages.pdf"},
            {},
            FakeArtifactLLM(),
            max_input_bytes=2000,
        )
        assert time.perf_counter() - start < 90
        assert document.recipe == recipe
        assert citations and all(
            item["chunk_id"] in {chunk.id for chunk in chunks} for item in citations
        )


@pytest.mark.asyncio
async def test_artifact_api_worker_selection_failure_retry_and_exports(
    level1_app: Any, monkeypatch: pytest.MonkeyPatch
) -> None:
    app, keys, _, engine = level1_app
    async with engine.begin() as connection:
        await connection.execute(text("PRAGMA foreign_keys=ON"))
    queue = QueueFake()

    async def make_pool(*_args: object, **_kwargs: object) -> QueueFake:
        return queue

    monkeypatch.setattr(artifacts_router, "create_pool", make_pool)
    settings = Settings(llm_provider="fake", embed_backend="fake")
    async with httpx.AsyncClient(
        transport=httpx.ASGITransport(app=app), base_url="http://test"
    ) as client:
        notebook_id = await seeded_book(app, client, keys, "The launch code is BLUEBIRD-42.")
        async with app.state.session_factory() as session:
            session.add(
                Source(
                    id="unselected",
                    notebook_id=notebook_id,
                    user_id="owner",
                    name="Secret.txt",
                    kind="txt",
                    status="ready",
                    selected=False,
                    progress=100,
                    metadata_json={},
                )
            )
            await session.flush()
            secret = sample_chunk(2, source_id="unselected", notebook_id=notebook_id)
            secret.text = "This private source says the answer is RED-99."
            secret.char_end = len(secret.text)
            session.add(secret)
            await session.commit()
        owner, other = auth(keys.token("owner")), auth(keys.token("other"))
        path = f"/notebooks/{notebook_id}/artifacts"
        recipes = await client.get(f"{path}/recipes", headers=owner)
        assert recipes.status_code == 200 and len(recipes.json()) == 8
        assert (await client.get(path, headers=other)).status_code == 404
        rejected = await client.post(
            path, headers=owner, json={"recipe": "study_guide", "source_ids": ["unselected"]}
        )
        assert rejected.status_code == 422
        created = await client.post(
            path,
            headers=owner,
            json={
                "recipe": "study_guide",
                "source_ids": ["evidence"],
                "language": "English",
                "length": "short",
                "focus": "launch",
            },
        )
        assert created.status_code == 202
        artifact_id = created.json()["id"]
        assert created.json()["source_ids"] == ["evidence"]
        assert (await client.get(f"{path}/{artifact_id}", headers=other)).status_code == 404
        assert queue.jobs[-1][0] == "generate_text_artifact"
        invalid = ScriptedLLM(["not json", "still not json"])
        ctx: dict[str, Any] = {
            "settings": settings,
            "session_factory": app.state.session_factory,
            "redis": queue,
            "artifact_llm": (invalid, "fake"),
        }
        await generate_text_artifact(ctx, artifact_id)
        failed = (await client.get(f"{path}/{artifact_id}", headers=owner)).json()
        assert failed["status"] == "failed" and "repair" in failed["error"]
        assert len(invalid.calls) == 2
        assert "RED-99" not in invalid.calls[0][1]["content"]
        retried = await client.post(f"{path}/{artifact_id}/regenerate", headers=owner)
        assert retried.status_code == 202 and retried.json()["status"] == "queued"
        ctx["artifact_llm"] = (FakeArtifactLLM(), "fake")
        await generate_text_artifact(ctx, artifact_id)
        ready = (await client.get(f"{path}/{artifact_id}", headers=owner)).json()
        assert ready["status"] == "ready" and ready["model_name"] == "fake"
        assert ready["citations"][0]["chunk_id"] == "fact"
        assert ready["citations"][0]["source_id"] == "evidence"
        assert "BLUEBIRD-42" in ready["output"]["sections"][0]["markdown"]
        assert (
            await client.get(f"{path}/{artifact_id}/export/md", headers=other)
        ).status_code == 404
        md = await client.get(f"{path}/{artifact_id}/export/md", headers=owner)
        assert md.status_code == 200 and "[S1] Evidence.txt" in md.text
        docx = await client.get(f"{path}/{artifact_id}/export/docx", headers=owner)
        assert docx.status_code == 200
        opened = Document(BytesIO(docx.content))
        assert "BLUEBIRD-42" in "\n".join(paragraph.text for paragraph in opened.paragraphs)
        assert (await client.delete(f"{path}/{artifact_id}", headers=owner)).status_code == 204
        assert (await client.get(f"{path}/{artifact_id}", headers=owner)).status_code == 404
