"""Level 2 extraction, offsets, SSRF, and idempotent worker checks."""

import asyncio
import json
import os
import zipfile
from collections.abc import AsyncIterator
from io import BytesIO
from typing import Any

import aiohttp
import httpx
import pytest
from docx import Document
from hypothesis import given
from hypothesis import strategies as st
from opennotebook_api.routers import sources as sources_router
from opennotebook_core.chunking import chunk_document
from opennotebook_core.config import Settings
from opennotebook_core.embeddings import FakeEmbedding, Kind
from opennotebook_core.extractors import (
    DocxExtractor,
    Extractor,
    MarkdownExtractor,
    PdfExtractor,
    PptxExtractor,
    TextExtractor,
    WebExtractor,
)
from opennotebook_core.models import Base, Chunk, Job, Notebook, Source
from opennotebook_core.safe_fetch import UnsafeURL, fetch_public_url, validate_url
from opennotebook_core.storage import MemoryStorage
from opennotebook_worker.main import ingest_source, reindex_notebook
from pptx import Presentation
from reportlab.pdfgen import canvas  # type: ignore[import-untyped]
from sqlalchemy import func, select, text
from sqlalchemy.ext.asyncio import async_sessionmaker, create_async_engine
from sqlalchemy.pool import StaticPool
from test_level1 import auth

pytest_plugins = ["test_level1"]


def pdf_bytes(text: str) -> bytes:
    stream = BytesIO()
    document = canvas.Canvas(stream)
    document.drawString(72, 720, text)
    document.save()
    return stream.getvalue()


@pytest.mark.slow
def test_real_embedding_ranks_relevant_sentence() -> None:
    if os.environ.get("RUN_SLOW_EMBED") != "1":
        pytest.skip("Set RUN_SLOW_EMBED=1 to run the real model")
    import numpy as np
    import torch
    from sentence_transformers import SentenceTransformer

    model = SentenceTransformer(
        "google/embeddinggemma-2",
        device="cpu",
        config_kwargs={"vision_config": None, "audio_config": None},
        model_kwargs={"torch_dtype": torch.float32},
    )
    assert next(model.parameters()).dtype == torch.float32
    query = model.encode(
        ["how do I reset my password"],
        prompt_name="SearchQuery",
        truncate_dim=768,
        normalize_embeddings=True,
    )
    docs = model.encode(
        [
            "You can reset your password using the account recovery link on the sign-in page.",
            "The weather forecast predicts a sunny afternoon.",
            "The railway station has four platforms.",
            "A chef prepares fresh pasta every morning.",
        ],
        prompt_name="Document",
        truncate_dim=768,
        normalize_embeddings=True,
    )
    scores = docs @ query[0]
    assert int(np.argmax(scores)) == 0
    assert np.isfinite(docs).all()
    assert np.allclose(np.linalg.norm(docs, axis=1), 1, atol=0.001)


def test_extractors_return_exact_blocks() -> None:
    docx = Document()
    docx.add_heading("Chapter One", 1)
    docx.add_paragraph("A source-grounded answer.")
    office = BytesIO()
    docx.save(office)
    slides = Presentation()
    slide = slides.slides.add_slide(slides.slide_layouts[1])
    slide.shapes.title.text = "Slide title"
    slide.placeholders[1].text = "Slide body."
    deck = BytesIO()
    slides.save(deck)
    fixtures: list[tuple[Extractor, bytes]] = [
        (PdfExtractor(), pdf_bytes("A PDF sentence.")),
        (DocxExtractor(), office.getvalue()),
        (PptxExtractor(), deck.getvalue()),
        (TextExtractor(), b"A plain text sentence."),
        (
            WebExtractor(),
            b"<html><body><article><h1>Title</h1><p>A substantial news article "
            b"sentence with enough readable text for extraction.</p></article></body></html>",
        ),
    ]
    for extractor, data in fixtures:
        document = extractor.extract(data)
        assert document.blocks
        for block in document.blocks:
            assert document.text[block.char_start : block.char_end] == block.text
    assert PdfExtractor().extract(fixtures[0][1]).page_count == 1
    assert PptxExtractor().extract(fixtures[2][1]).page_count == 1
    markdown = MarkdownExtractor().extract(b"# Findings\n\nEvidence supports the result.")
    assert markdown.blocks[-1].heading_path == ("Findings",)


@given(
    st.lists(
        st.text(
            alphabet=st.characters(whitelist_categories=("L", "N")), min_size=1, max_size=30
        ).filter(str.strip),
        min_size=1,
        max_size=30,
    )
)
def test_chunk_offsets_and_sentence_boundaries(sentences: list[str]) -> None:
    value = " ".join(f"{sentence}." for sentence in sentences)
    document = TextExtractor().extract(value.encode())
    chunks = chunk_document(document, target_tokens=8, overlap=0.12)
    assert chunks
    for chunk in chunks:
        assert document.text[chunk.char_start : chunk.char_end] == chunk.text
        assert chunk.text.endswith(".")
    assert chunks[0].char_start == 0
    assert chunks[-1].char_end == len(value)
    for left, right in zip(chunks, chunks[1:], strict=False):
        assert left.char_end >= right.char_start


def test_abbreviations_do_not_end_chunks() -> None:
    document = TextExtractor().extract(
        b"Dr. Smith documented the result. The second sentence verifies it."
    )
    chunks = chunk_document(document, target_tokens=8)
    assert chunks[0].text.startswith("Dr. Smith")
    assert all(document.text[chunk.char_start : chunk.char_end] == chunk.text for chunk in chunks)


def test_chunks_overlap_on_sentence_boundaries() -> None:
    document = TextExtractor().extract(
        " ".join(f"Sentence {number} gives evidence." for number in range(20)).encode()
    )
    chunks = chunk_document(document, target_tokens=20, overlap=0.12)
    assert len(chunks) > 2
    assert all(
        left.char_end > right.char_start for left, right in zip(chunks, chunks[1:], strict=False)
    )


@pytest.mark.parametrize(
    "url",
    [
        "http://127.0.0.1/",
        "http://10.1.2.3/",
        "http://169.254.169.254/latest/meta-data/",
        "http://localhost/admin",
        "http://[::1]/",
        "file:///etc/passwd",
    ],
)
def test_ssrf_rejects_private_and_unsafe_urls(url: str) -> None:
    with pytest.raises(UnsafeURL):
        validate_url(url)


@pytest.mark.asyncio
async def test_fetch_rejects_redirect_to_private_and_oversize_and_timeout(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    class Response:
        def __init__(
            self,
            status: int,
            headers: dict[str, str] | None = None,
            size: int | None = None,
            slow: bool = False,
        ):
            self.status = status
            self.headers = headers or {}
            self.content_length = size
            self.slow = slow
            self.content = self

        async def __aenter__(self) -> "Response":
            return self

        async def __aexit__(self, *_: Any) -> None:
            return None

        def raise_for_status(self) -> None:
            return None

        async def iter_chunked(self, _: int) -> AsyncIterator[bytes]:
            if self.slow:
                raise TimeoutError("slow response")
            yield b"x" * (self.content_length or 1)

    class Session:
        def __init__(self, response: Response):
            self.response = response

        async def __aenter__(self) -> "Session":
            return self

        async def __aexit__(self, *_: Any) -> None:
            return None

        def get(self, *_: Any, **__: Any) -> Response:
            return self.response

    def serve(response: Response) -> None:
        monkeypatch.setattr(aiohttp, "ClientSession", lambda **_: Session(response))

    serve(Response(302, {"Location": "http://169.254.169.254/latest/meta-data/"}))
    with pytest.raises(UnsafeURL):
        await fetch_public_url("https://public.example/article")
    serve(Response(200, size=101))
    with pytest.raises(ValueError, match="size limit"):
        await fetch_public_url("https://public.example/article", max_bytes=100)
    serve(Response(200, slow=True))
    with pytest.raises(TimeoutError, match="slow response"):
        await fetch_public_url("https://public.example/article")


class FakeRedis:
    def __init__(self) -> None:
        self.messages: list[str] = []

    async def publish(self, channel: str, message: str) -> int:
        self.messages.append(message)
        return 1


@pytest.mark.asyncio
async def test_worker_idempotency_failure_and_retry() -> None:
    engine = create_async_engine("sqlite+aiosqlite:///:memory:", poolclass=StaticPool)
    async with engine.begin() as connection:
        await connection.run_sync(Base.metadata.create_all)
    factory = async_sessionmaker(engine, expire_on_commit=False)
    storage = MemoryStorage()
    redis = FakeRedis()
    settings = Settings(embed_backend="fake", storage_backend="memory")
    context: dict[str, Any] = {
        "session_factory": factory,
        "storage": storage,
        "redis": redis,
        "settings": settings,
        "embedding": FakeEmbedding(),
    }
    async with factory() as session:
        session.add(Notebook(id="book", user_id="owner", title="Book", response_length="medium"))
        session.add(
            Source(
                id="source",
                notebook_id="book",
                user_id="owner",
                name="report.pdf",
                kind="pdf",
                storage_key="report.pdf",
                status="queued",
                progress=0,
                metadata_json={},
            )
        )
        session.add(Job(id="job1", user_id="owner", kind="ingest_source", status="queued"))
        await session.commit()
    await storage.put("report.pdf", b"%PDF-corrupt")
    await ingest_source(context, "source", "job1")
    async with factory() as session:
        source = await session.get(Source, "source")
        assert source and source.status == "failed" and "PDF" in (source.error or "")
        assert await session.scalar(select(func.count(Chunk.id))) == 0
        session.add(Job(id="job2", user_id="owner", kind="ingest_source", status="queued"))
        await session.commit()
    await storage.put("report.pdf", pdf_bytes("A valid research sentence."))
    await ingest_source(context, "source", "job2")
    await ingest_source(context, "source", "job2")
    async with factory() as session:
        source = await session.get(Source, "source")
        chunks = list(
            (await session.scalars(select(Chunk).where(Chunk.source_id == "source"))).all()
        )
        assert source and source.status == "ready" and source.extracted_text
        assert len(chunks) == 1
        assert source.extracted_text[chunks[0].char_start : chunks[0].char_end] == chunks[0].text
        assert source.embedding_model == "fake" and source.embedding_dim == 768
    stages = [json.loads(message)["stage"] for message in redis.messages]
    stage_transitions = [
        stage for index, stage in enumerate(stages) if index == 0 or stage != stages[index - 1]
    ]
    assert stage_transitions[-6:] == [
        "fetching",
        "extracting",
        "chunking",
        "embedding",
        "indexing",
        "ready",
    ]
    await engine.dispose()


@pytest.mark.asyncio
async def test_two_users_ingest_concurrently() -> None:
    engine = create_async_engine("sqlite+aiosqlite:///:memory:", poolclass=StaticPool)
    async with engine.begin() as connection:
        await connection.run_sync(Base.metadata.create_all)
    factory = async_sessionmaker(engine, expire_on_commit=False)
    storage = MemoryStorage()

    class SlowEmbedding(FakeEmbedding):
        active = 0
        peak = 0

        async def embed(
            self, texts: list[str], kind: Kind, dim: int | None = None
        ) -> list[list[float]]:
            self.active += 1
            self.peak = max(self.peak, self.active)
            await asyncio.sleep(0.05)
            try:
                return await super().embed(texts, kind, dim)
            finally:
                self.active -= 1

    embedding = SlowEmbedding()
    context: dict[str, Any] = {
        "session_factory": factory,
        "storage": storage,
        "redis": FakeRedis(),
        "settings": Settings(embed_backend="fake", storage_backend="memory"),
        "embedding": embedding,
    }
    async with factory() as session:
        for owner in ("alice", "bob"):
            session.add(
                Notebook(id=f"book-{owner}", user_id=owner, title=owner, response_length="medium")
            )
            session.add(
                Source(
                    id=f"source-{owner}",
                    notebook_id=f"book-{owner}",
                    user_id=owner,
                    name="notes.txt",
                    kind="txt",
                    storage_key=f"{owner}.txt",
                    status="queued",
                    progress=0,
                    metadata_json={},
                )
            )
            session.add(
                Job(id=f"job-{owner}", user_id=owner, kind="ingest_source", status="queued")
            )
            await storage.put(f"{owner}.txt", f"{owner} research sentence.".encode())
        await session.commit()
    await asyncio.gather(
        ingest_source(context, "source-alice", "job-alice"),
        ingest_source(context, "source-bob", "job-bob"),
    )
    async with factory() as session:
        sources = (await session.scalars(select(Source))).all()
        assert len(sources) == 2
        assert all(source.status == "ready" for source in sources), [
            (source.user_id, source.error) for source in sources
        ]
        assert await session.scalar(select(func.count(Chunk.id))) == 2
    assert embedding.peak == 2
    await engine.dispose()


@pytest.mark.asyncio
async def test_source_and_events_ownership(level1_app: Any) -> None:
    app, keys, _, _ = level1_app
    async with httpx.AsyncClient(
        transport=httpx.ASGITransport(app=app), base_url="http://test"
    ) as client:
        created = await client.post(
            "/notebooks", headers=auth(keys.token("owner")), json={"title": "Private"}
        )
        assert created.status_code == 201
        notebook_id = created.json()["id"]
        owner = auth(keys.token("owner"))
        other = auth(keys.token("other"))
        assert (
            await client.get(f"/notebooks/{notebook_id}/sources", headers=other)
        ).status_code == 404
        assert (
            await client.get(f"/notebooks/{notebook_id}/events", headers=other)
        ).status_code == 404
        assert (await client.get(f"/notebooks/{notebook_id}/sources", headers=owner)).json() == []


@pytest.mark.asyncio
async def test_deleting_source_removes_file_and_chunks(
    level1_app: Any, monkeypatch: pytest.MonkeyPatch
) -> None:
    app, keys, _, _ = level1_app
    storage = MemoryStorage()
    monkeypatch.setattr(sources_router, "make_storage", lambda _: storage)
    async with httpx.AsyncClient(
        transport=httpx.ASGITransport(app=app), base_url="http://test"
    ) as client:
        owner = auth(keys.token("owner"))
        other = auth(keys.token("other"))
        created = await client.post("/notebooks", headers=owner, json={"title": "Delete test"})
        notebook_id = created.json()["id"]
        await storage.put("owned.pdf", b"%PDF-source")
        async with app.state.session_factory() as session:
            await session.execute(text("PRAGMA foreign_keys=ON"))
            session.add(
                Source(
                    id="owned-source",
                    notebook_id=notebook_id,
                    user_id="owner",
                    name="owned.pdf",
                    kind="pdf",
                    storage_key="owned.pdf",
                    status="ready",
                    progress=100,
                    embedding_model="fake",
                    embedding_dim=768,
                    metadata_json={},
                )
            )
            await session.flush()
            session.add(
                Chunk(
                    id="owned-chunk",
                    source_id="owned-source",
                    notebook_id=notebook_id,
                    ordinal=0,
                    text="Source text.",
                    char_start=0,
                    char_end=12,
                    heading_path=[],
                    embedding=[1.0] + [0.0] * 767,
                    metadata_json={},
                )
            )
            await session.commit()
        path = f"/notebooks/{notebook_id}/sources/owned-source"
        assert (await client.delete(path, headers=other)).status_code == 404
        assert (await client.delete(path, headers=owner)).status_code == 204
        async with app.state.session_factory() as session:
            assert await session.get(Source, "owned-source") is None
            assert await session.scalar(select(func.count(Chunk.id))) == 0
        assert "owned.pdf" not in storage.items


@pytest.mark.asyncio
async def test_upload_rejects_oversize_and_wrong_file_content(level1_app: Any) -> None:
    app, keys, _, _ = level1_app
    app.state.settings.upload_max_mb = 1
    async with httpx.AsyncClient(
        transport=httpx.ASGITransport(app=app), base_url="http://test"
    ) as client:
        headers = auth(keys.token("owner"))
        created = await client.post("/notebooks", headers=headers, json={"title": "Limits"})
        notebook_id = created.json()["id"]
        path = f"/notebooks/{notebook_id}/sources/upload"
        too_big = await client.post(
            path,
            headers=headers,
            files={"file": ("oversize.txt", b"x" * (1024 * 1024 + 1), "text/plain")},
        )
        assert too_big.status_code == 413
        wrong_pdf = await client.post(
            path,
            headers=headers,
            files={"file": ("fake.pdf", b"not a PDF", "application/pdf")},
        )
        assert wrong_pdf.status_code == 415
        wrong_office = await client.post(
            path,
            headers=headers,
            files={
                "file": (
                    "fake.docx",
                    b"PK-not-a-zip",
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                )
            },
        )
        assert wrong_office.status_code == 415

        zip_bomb = BytesIO()
        with zipfile.ZipFile(zip_bomb, "w", compression=zipfile.ZIP_DEFLATED) as archive:
            archive.writestr("word/document.xml", "A" * (201 * 1024 * 1024))
        rejected_bomb = await client.post(
            path,
            headers=headers,
            files={
                "file": (
                    "expanded.docx",
                    zip_bomb.getvalue(),
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                )
            },
        )
        assert rejected_bomb.status_code == 413


@pytest.mark.asyncio
async def test_all_non_pdf_sources_ingest_and_reindex() -> None:
    engine = create_async_engine("sqlite+aiosqlite:///:memory:", poolclass=StaticPool)
    async with engine.begin() as connection:
        await connection.run_sync(Base.metadata.create_all)
    factory = async_sessionmaker(engine, expire_on_commit=False)
    storage = MemoryStorage()
    document = Document()
    document.add_paragraph("A DOCX paragraph about research.")
    docx_stream = BytesIO()
    document.save(docx_stream)
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[1])
    slide.shapes.title.text = "Research presentation"
    slide.placeholders[1].text = "A PPTX paragraph about research."
    pptx_stream = BytesIO()
    deck.save(pptx_stream)
    files = {
        "docx": docx_stream.getvalue(),
        "pptx": pptx_stream.getvalue(),
        "txt": b"A TXT paragraph about research.",
        "md": b"# Research\n\nA Markdown paragraph about research.",
        "text": b"A pasted paragraph about research.",
    }

    async def fake_fetch(_: str) -> bytes:
        return (
            b"<html><body><article><h1>News</h1><p>A substantial article about "
            b"research discoveries and the methods used to establish them."
            b"</p></article></body></html>"
        )

    settings = Settings(embed_backend="fake", storage_backend="memory")
    context: dict[str, Any] = {
        "session_factory": factory,
        "storage": storage,
        "redis": FakeRedis(),
        "settings": settings,
        "embedding": FakeEmbedding(),
        "fetcher": fake_fetch,
    }
    async with factory() as session:
        session.add(Notebook(id="book", user_id="owner", title="Book", response_length="medium"))
        for kind in (*files, "web"):
            session.add(
                Source(
                    id=kind,
                    notebook_id="book",
                    user_id="owner",
                    name=kind,
                    kind=kind,
                    storage_key=kind if kind != "web" else None,
                    url="https://public.example/article" if kind == "web" else None,
                    status="queued",
                    progress=0,
                    metadata_json={},
                )
            )
            session.add(
                Job(id=f"job-{kind}", user_id="owner", kind="ingest_source", status="queued")
            )
        await session.commit()
    for kind, data in files.items():
        await storage.put(kind, data)
    for kind in (*files, "web"):
        await ingest_source(context, kind, f"job-{kind}")
    async with factory() as session:
        sources = list(
            (await session.scalars(select(Source).where(Source.notebook_id == "book"))).all()
        )
        assert len(sources) == 6
        assert all(source.status == "ready" for source in sources), [
            (source.kind, source.error) for source in sources
        ]
        count = await session.scalar(select(func.count(Chunk.id)))
        assert count is not None and count >= 6
        sources[0].embedding_model = "previous-model"
        session.add(
            Source(
                id="new",
                notebook_id="book",
                user_id="owner",
                name="new",
                kind="txt",
                storage_key="new",
                status="queued",
                progress=0,
                metadata_json={},
            )
        )
        session.add(Job(id="job-new", user_id="owner", kind="ingest_source", status="queued"))
        session.add(
            Job(id="reindex-job", user_id="owner", kind="reindex_notebook", status="queued")
        )
        await session.commit()
    await storage.put("new", b"A new source that uses the current model.")
    await ingest_source(context, "new", "job-new")
    async with factory() as session:
        new_source = await session.get(Source, "new")
        assert new_source and new_source.status == "failed"
        assert "Re-index" in (new_source.error or "")
        assert (
            await session.scalar(select(func.count(Chunk.id)).where(Chunk.source_id == "new")) == 0
        )
    await reindex_notebook(context, "book", "reindex-job")
    async with factory() as session:
        sources = list(
            (await session.scalars(select(Source).where(Source.notebook_id == "book"))).all()
        )
        assert all(
            source.embedding_model == "fake" and source.embedding_dim == 768
            for source in sources
            if source.id != "new"
        )
        job = await session.get(Job, "reindex-job")
        assert job and job.status == "succeeded"
        session.add(Job(id="retry-new", user_id="owner", kind="ingest_source", status="queued"))
        await session.commit()
    await ingest_source(context, "new", "retry-new")
    async with factory() as session:
        new_source = await session.get(Source, "new")
        assert new_source and new_source.status == "ready" and new_source.embedding_model == "fake"
    await engine.dispose()
