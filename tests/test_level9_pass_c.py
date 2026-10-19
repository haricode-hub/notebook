"""Level 9 security, export, and multimodal ranking acceptance tests."""

import base64
import io
import zipfile
from collections import defaultdict
from typing import Any, cast

import httpx
import pytest
from opennotebook_core.embeddings import EmbeddingProvider
from opennotebook_core.extractors import extract_figures
from opennotebook_core.models import Artifact, Chunk, MediaItem, Note, Source
from opennotebook_core.prompts.chat_v1 import build_chat_messages, validate_citations
from opennotebook_core.retrieval import rrf_merge_modalities, select_context
from test_level1 import auth
from test_level3_chat import seeded_book

pytest_plugins = ["test_level1"]


class _Redis:
    def __init__(self) -> None:
        self.counts: defaultdict[str, int] = defaultdict(int)

    async def eval(self, _script: str, _keys: int, key: str) -> int:
        self.counts[key] += 1
        return self.counts[key]

    async def aclose(self) -> None:
        return None


class _QueryEmbedding:
    async def embed(self, texts: list[str], kind: str, dim: int) -> list[list[float]]:
        del kind
        vectors = []
        for text in texts:
            vector = [0.0] * dim
            vector[0 if text.casefold().startswith("chart") else 1] = 1.0
            vectors.append(vector)
        return vectors


@pytest.mark.asyncio
async def test_share_revocation_read_only_scope_and_zip_contents(
    level1_app: Any, monkeypatch: pytest.MonkeyPatch
) -> None:
    from opennotebook_api.routers import sharing

    app, keys, _, _ = level1_app
    redis = _Redis()
    redis_type = cast(Any, sharing).Redis
    monkeypatch.setattr(redis_type, "from_url", lambda *_args, **_kwargs: redis)
    async with httpx.AsyncClient(
        transport=httpx.ASGITransport(app=app), base_url="http://test"
    ) as client:
        notebook_id = await seeded_book(app, client, keys, "The chart shows a steady increase.")
        owner, other = auth(keys.token("owner")), auth(keys.token("other"))
        async with app.state.session_factory() as db:
            source = await db.get(Source, "evidence")
            assert source is not None
            source.extracted_text = "The chart shows a steady increase."
            db.add(
                Note(
                    notebook_id=notebook_id,
                    user_id="owner",
                    title="Shared note",
                    content_markdown="A sourced note.",
                    citations_json=[],
                )
            )
            db.add(
                Artifact(
                    notebook_id=notebook_id,
                    user_id="owner",
                    recipe="briefing",
                    status="ready",
                    stage="ready",
                    progress=100,
                    options_json={},
                    source_ids=["evidence"],
                    output_json={"title": "Briefing", "sections": []},
                    citations_json=[],
                )
            )
            await db.commit()

        created = await client.post(f"/notebooks/{notebook_id}/shares", headers=owner, json={})
        assert created.status_code == 201
        token, share_id = created.json()["token"], created.json()["id"]
        public = await client.get(f"/public/shares/{token}")
        assert public.status_code == 200
        assert public.json()["sources"][0]["text"] == "The chart shows a steady increase."
        assert (
            await client.post(f"/public/shares/{token}", json={"title": "mutate"})
        ).status_code == 405
        assert (
            await client.post(f"/notebooks/{notebook_id}/shares", headers=other, json={})
        ).status_code == 404
        assert (
            await client.get(f"/notebooks/{notebook_id}/export.zip", headers=other)
        ).status_code == 404

        exported = await client.get(f"/notebooks/{notebook_id}/export.zip", headers=owner)
        assert exported.status_code == 200
        with zipfile.ZipFile(io.BytesIO(exported.content)) as archive:
            names = set(archive.namelist())
            assert {"notebook.json", "sources.json"} <= names
            note_markdown = next(
                name for name in names if name.startswith("notes/") and name.endswith(".md")
            )
            assert archive.read(note_markdown).decode() == "A sourced note."
            artifact_file = next(
                name
                for name in names
                if name.startswith("artifacts/") and name.endswith(".json")
            )
            assert b'"recipe": "briefing"' in archive.read(artifact_file)

        assert (
            await client.delete(f"/notebooks/{notebook_id}/shares/{share_id}", headers=owner)
        ).status_code == 204
        assert (await client.get(f"/public/shares/{token}")).status_code == 404


def test_pptx_figure_extractor_keeps_slide_and_bounds() -> None:
    from PIL import Image
    from pptx import Presentation
    from pptx.util import Inches

    png = io.BytesIO()
    Image.new("RGB", (20, 20), (25, 60, 120)).save(png, format="PNG")
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.shapes.add_textbox(
        Inches(1), Inches(0.3), Inches(8), Inches(0.4)
    ).text = "Chart: revenue rises each quarter"
    slide.shapes.add_picture(io.BytesIO(png.getvalue()), Inches(1), Inches(1), width=Inches(2))
    output = io.BytesIO()
    presentation.save(output)

    figures = extract_figures("pptx", output.getvalue())
    assert len(figures) == 1
    assert figures[0].page == 1
    assert figures[0].bbox is not None and figures[0].bbox["width"] > 0
    assert "revenue rises" in figures[0].text_anchor


def test_multimodal_rrf_uses_ranks_not_cross_modal_scores() -> None:
    text = Chunk(
        id="z-text",
        source_id="source",
        notebook_id="book",
        ordinal=0,
        text="fact",
        char_start=0,
        char_end=4,
        heading_path=[],
        embedding=[1.0],
        metadata_json={},
    )
    image = MediaItem(
        id="a-image",
        source_id="source",
        notebook_id="book",
        storage_key="figure.png",
        text_anchor="chart",
        page=1,
        modality="image",
        embedding=[1.0],
    )
    cast(Any, text).raw_similarity = -10000.0
    cast(Any, image).raw_similarity = 10000.0
    # Rank-one ties use stable IDs, regardless of unrelated score attributes.
    assert [item.id for item in rrf_merge_modalities([text], [image])] == ["a-image", "z-text"]


@pytest.mark.asyncio
async def test_chart_caption_retrieves_and_validates_figure_citation(level1_app: Any) -> None:
    app, keys, _, engine = level1_app
    async with httpx.AsyncClient(
        transport=httpx.ASGITransport(app=app), base_url="http://test"
    ) as client:
        notebook_id = await seeded_book(
            app, client, keys, "Quarterly revenue is shown in the chart."
        )
    vector = [0.0] * 768
    vector[0] = 1.0
    async with app.state.session_factory() as db:
        # Align text with the deterministic query; image ranking has an irrelevant candidate.
        source = await db.get(Source, "evidence")
        assert source is not None
        source.selected = True
        chunk = await db.get(Chunk, "fact")
        assert chunk is not None
        chunk.embedding = vector
        db.add_all(
            [
                MediaItem(
                    id="chart",
                    source_id="evidence",
                    notebook_id=notebook_id,
                    storage_key="media/chart.png",
                    text_anchor="Chart: revenue rises every quarter",
                    page=2,
                    modality="image",
                    embedding=vector,
                ),
                MediaItem(
                    id="other",
                    source_id="evidence",
                    notebook_id=notebook_id,
                    storage_key="media/other.png",
                    text_anchor="A map of mountain ranges",
                    page=3,
                    modality="image",
                    embedding=[-1.0, *([0.0] * 767)],
                ),
            ]
        )
        await db.commit()
        selection = await select_context(
            db,
            cast(EmbeddingProvider, _QueryEmbedding()),
            notebook_id,
            ["evidence"],
            "chart quarterly revenue",
            multimodal_enabled=True,
        )
    assert any(item.id == "chart" for item in (selection.media_items or []))
    labels = {
        number: item
        for number, item in enumerate([*selection.chunks, *(selection.media_items or [])], start=1)
    }
    figure_label = next(
        number
        for number, item in labels.items()
        if isinstance(item, MediaItem) and item.id == "chart"
    )
    prompt = build_chat_messages(
        selection, "What does the chart show?", [], persona=None, response_length="short"
    )
    assert "Chart: revenue rises" in prompt[-1]["content"]
    answer, citations = validate_citations(
        f"Revenue rises by quarter [S{figure_label}].", selection
    )
    assert "Revenue rises" in answer
    assert citations[0]["media_item_id"] == "chart"
    await engine.dispose()


@pytest.mark.asyncio
async def test_embed_media_is_flagged_and_fake_backend_does_not_claim_support() -> None:
    from opennotebook_core.config import Settings
    from opennotebook_embed.main import create_app

    payload = {
        "items": [
            {
                "modality": "image",
                "media_base64": base64.b64encode(b"valid-enough-payload").decode(),
                "text": "chart",
            }
        ],
        "dim": 768,
    }
    for enabled, expected in ((False, 404), (True, 501)):
        app = create_app(Settings(embed_backend="fake", multimodal_enabled=enabled))
        async with httpx.AsyncClient(
            transport=httpx.ASGITransport(app=app), base_url="http://embed"
        ) as client:
            response = await client.post("/embed-media", json=payload)
        assert response.status_code == expected
