"""Reusable source-grounded artifact recipe registry and generator."""

import asyncio
import html
import json
import re
from collections.abc import AsyncIterator, Awaitable, Callable
from dataclasses import dataclass
from io import BytesIO
from typing import Literal, Protocol, cast

from docx import Document
from pydantic import BaseModel, ConfigDict, Field, model_validator

from opennotebook_core.audio import AudioOverview
from opennotebook_core.models import Chunk
from opennotebook_core.prompts.artifact_v2 import RECIPES, artifact_messages, map_messages
from opennotebook_core.prompts.chat_v1 import MARKER, Citation


class ArtifactSection(BaseModel):
    heading: str = Field(min_length=1, max_length=200)
    markdown: str = Field(min_length=1)
    citation_refs: list[int] = Field(min_length=1)


class ArtifactDocument(BaseModel):
    model_config = ConfigDict(extra="forbid")
    recipe: str
    title: str = Field(min_length=1, max_length=200)
    sections: list[ArtifactSection] = Field(min_length=1)


class MindMapNode(BaseModel):
    model_config = ConfigDict(extra="forbid", strict=True)
    label: str = Field(min_length=1, max_length=120)
    summary: str = Field(min_length=1, max_length=1000)
    source_refs: list[int] = Field(min_length=1, max_length=10)
    children: list["MindMapNode"] = Field(default_factory=list, max_length=12)

    @model_validator(mode="after")
    def valid_source_refs(self) -> "MindMapNode":
        if any(ref < 1 for ref in self.source_refs) or len(set(self.source_refs)) != len(
            self.source_refs
        ):
            raise ValueError("Mind map source_refs must be distinct positive citation numbers")
        return self


class MindMapArtifact(BaseModel):
    model_config = ConfigDict(extra="forbid", strict=True)
    recipe: Literal["mindmap"] = "mindmap"
    title: str = Field(min_length=1, max_length=200)
    root: MindMapNode

    @model_validator(mode="after")
    def bounded_tree(self) -> "MindMapArtifact":
        count = 0

        def walk(node: MindMapNode, depth: int) -> None:
            nonlocal count
            count += 1
            if depth > 4:
                raise ValueError("Mind map depth must be at most 4")
            if count > 60:
                raise ValueError("Mind map must contain at most 60 nodes")
            for child in node.children:
                walk(child, depth + 1)

        walk(self.root, 1)
        return self


class Flashcard(BaseModel):
    model_config = ConfigDict(extra="forbid", strict=True)
    front: str = Field(min_length=1, max_length=500)
    back: str = Field(min_length=1, max_length=1500)
    source_ref: int = Field(gt=0)


class FlashcardsArtifact(BaseModel):
    model_config = ConfigDict(extra="forbid", strict=True)
    recipe: Literal["flashcards"] = "flashcards"
    title: str = Field(min_length=1, max_length=200)
    cards: list[Flashcard] = Field(min_length=1, max_length=60)


class QuizQuestion(BaseModel):
    model_config = ConfigDict(extra="forbid", strict=True)
    question: str = Field(min_length=1, max_length=800)
    options: list[str] = Field(min_length=4, max_length=4)
    answer_index: int = Field(ge=0, le=3)
    explanation: str = Field(min_length=1, max_length=1500)
    source_ref: int = Field(gt=0)

    @model_validator(mode="after")
    def distinct_options_and_cited_explanation(self) -> "QuizQuestion":
        if len({option.strip().casefold() for option in self.options}) != 4:
            raise ValueError("Quiz options must be distinct")
        refs = {int(match.group(1)) for match in MARKER.finditer(self.explanation)}
        if refs != {self.source_ref}:
            raise ValueError("Quiz explanation must cite its source_ref")
        return self


class QuizArtifact(BaseModel):
    model_config = ConfigDict(extra="forbid", strict=True)
    recipe: Literal["quiz"] = "quiz"
    title: str = Field(min_length=1, max_length=200)
    questions: list[QuizQuestion] = Field(min_length=1, max_length=50)


class EvidenceItem(BaseModel):
    ref: int
    fact: str = Field(min_length=1, max_length=800)


class EvidenceDigest(BaseModel):
    items: list[EvidenceItem] = Field(min_length=1, max_length=3)


@dataclass(frozen=True)
class ArtifactRecipe:
    id: str
    title: str
    instruction: str
    renderer_id: str = "markdown_sections_v1"
    output_schema: type[BaseModel] = ArtifactDocument


RECIPES_REGISTRY: dict[str, ArtifactRecipe] = {
    key: ArtifactRecipe(key, title, instruction) for key, (title, instruction) in RECIPES.items()
}
RECIPES_REGISTRY["audio_overview"] = ArtifactRecipe(
    "audio_overview",
    "Audio overview",
    "Create a grounded two-host conversation.",
    "audio_player_v1",
    AudioOverview,
)
for _id, _schema, _title, _renderer in (
    ("mindmap", MindMapArtifact, "Mind map", "mindmap_v1"),
    ("flashcards", FlashcardsArtifact, "Flashcards", "flashcards_v1"),
    ("quiz", QuizArtifact, "Quiz", "quiz_v1"),
):
    RECIPES_REGISTRY[_id] = ArtifactRecipe(
        _id,
        _title,
        RECIPES[_id][1],
        _renderer,
        _schema,
    )

GeneratedArtifact = (
    ArtifactDocument | MindMapArtifact | FlashcardsArtifact | QuizArtifact | AudioOverview
)


class ArtifactLLM(Protocol):
    def stream(self, messages: list[dict[str, str]]) -> AsyncIterator[str]: ...


async def _collect(llm: ArtifactLLM, messages: list[dict[str, str]]) -> str:
    json_stream = getattr(llm, "stream_json", None)
    stream = json_stream(messages) if callable(json_stream) else llm.stream(messages)
    return "".join([piece async for piece in stream])


def _extract_json_object(raw: str) -> str:
    """Accept a JSON object wrapped in fences or short model preambles."""
    cleaned = re.sub(r"^\s*```(?:json)?\s*|\s*```\s*$", "", raw.strip(), flags=re.I)
    start = cleaned.find("{")
    if start < 0:
        return cleaned
    depth = 0
    quoted = False
    escaped = False
    for index in range(start, len(cleaned)):
        char = cleaned[index]
        if quoted:
            if escaped:
                escaped = False
            elif char == "\\":
                escaped = True
            elif char == '"':
                quoted = False
        elif char == '"':
            quoted = True
        elif char == "{":
            depth += 1
        elif char == "}":
            depth -= 1
            if depth == 0:
                return cleaned[start : index + 1]
    return cleaned


async def _validated_call[T: BaseModel](
    llm: ArtifactLLM,
    messages: list[dict[str, str]],
    schema: type[T],
    validate: Callable[[T], None],
) -> T:
    schema_json = json.dumps(schema.model_json_schema(), ensure_ascii=False)
    messages = [dict(message) for message in messages]
    messages[0]["content"] += (
        "\nReturn one JSON object that validates against this exact JSON Schema. "
        f"Do not add fields or wrap the object. Schema: {schema_json}"
    )
    raw = await _collect(llm, messages)
    for attempt in range(2):
        try:
            cleaned = _extract_json_object(raw)
            value = schema.model_validate_json(cleaned)
            validate(value)
            return value
        except ValueError as exc:
            if attempt:
                raise ValueError(
                    "Artifact model returned invalid or ungrounded JSON after repair"
                ) from exc
            messages = [
                *messages,
                {"role": "assistant", "content": raw[:12000]},
                {
                    "role": "user",
                    "content": (
                        "Repair the JSON to match the schema and citation rules. "
                        f"Validation error: {str(exc)[:1000]}. Return JSON only and follow "
                        f"this exact JSON Schema: {schema_json}"
                    ),
                },
            ]
            raw = await _collect(llm, messages)
    raise AssertionError("unreachable")


def _batches(excerpts: list[dict[str, object]], max_bytes: int) -> list[list[dict[str, object]]]:
    batches: list[list[dict[str, object]]] = []
    batch: list[dict[str, object]] = []
    for excerpt in excerpts:
        item = dict(excerpt)
        # A single very large chunk still has a strict request budget.
        item["text"] = str(item["text"])[: max_bytes // 2]
        if batch and len(json.dumps([*batch, item], ensure_ascii=False).encode()) > max_bytes:
            batches.append(batch)
            batch = []
        batch.append(item)
    if batch:
        batches.append(batch)
    return batches


def _validate_document(document: ArtifactDocument, recipe: str, allowed: set[int]) -> None:
    if document.recipe != recipe:
        raise ValueError("Wrong recipe in artifact output")
    for section in document.sections:
        markers = {int(match.group(1)) for match in MARKER.finditer(section.markdown)}
        if not markers or markers != set(section.citation_refs) or not markers <= allowed:
            raise ValueError("Artifact section contains missing or fabricated citations")


def _validate_output(
    document: BaseModel, recipe: str, allowed: set[int], options: dict[str, object]
) -> None:
    if getattr(document, "recipe", None) != recipe:
        raise ValueError("Wrong recipe in artifact output")
    if isinstance(document, ArtifactDocument):
        _validate_document(document, recipe, allowed)
        return
    refs: set[int] = set()
    if isinstance(document, MindMapArtifact):

        def walk(node: MindMapNode) -> None:
            refs.update(node.source_refs)
            for child in node.children:
                walk(child)

        walk(document.root)
    elif isinstance(document, FlashcardsArtifact):
        refs.update(card.source_ref for card in document.cards)
    elif isinstance(document, QuizArtifact):
        raw_count = options.get("question_count", 10)
        if not isinstance(raw_count, int):
            raise ValueError("Quiz question count must be an integer")
        requested = raw_count
        if len(document.questions) != requested:
            raise ValueError("Quiz question count must match the requested count")
        refs.update(question.source_ref for question in document.questions)
    if not refs or not refs <= allowed:
        raise ValueError("Artifact contains missing or fabricated citations")


async def generate_artifact(
    recipe: str,
    chunks: list[Chunk],
    source_names: dict[str, str],
    options: dict[str, object],
    llm: ArtifactLLM,
    *,
    max_input_bytes: int = 10000,
    progress: Callable[[str, int], Awaitable[None]] | None = None,
) -> tuple[GeneratedArtifact, list[Citation]]:
    """Map long sources into cited evidence, then generate a validated document."""
    if recipe not in RECIPES_REGISTRY:
        raise ValueError("Unknown artifact recipe")
    if not chunks:
        raise ValueError("Select at least one ready source with indexed text")
    labeled = {n: chunk for n, chunk in enumerate(chunks, 1)}
    excerpts: list[dict[str, object]] = [
        {
            "ref": n,
            "source": source_names.get(chunk.source_id, "Source"),
            "page": chunk.page_start,
            "text": chunk.text,
        }
        for n, chunk in labeled.items()
    ]
    batches = _batches(excerpts, max_input_bytes)
    if len(batches) > 1:
        semaphore = asyncio.Semaphore(4)

        async def map_batch(
            batch: list[dict[str, object]], reducing: bool = False
        ) -> EvidenceDigest:
            allowed = {int(str(item["ref"])) for item in batch}

            def validate(digest: EvidenceDigest) -> None:
                if any(item.ref not in allowed for item in digest.items):
                    raise ValueError("Evidence map invented a citation")

            async with semaphore:
                return await _validated_call(
                    llm, map_messages(batch, reducing=reducing), EvidenceDigest, validate
                )

        digests = await asyncio.gather(*(map_batch(batch) for batch in batches))
        items = [item.model_dump() for digest in digests for item in digest.items]
        if progress:
            await progress("reducing", 65)
        for _ in range(10):
            if len(json.dumps(items, ensure_ascii=False).encode()) <= max_input_bytes:
                break
            reduction_input: list[dict[str, object]] = [
                {"ref": item["ref"], "text": str(item["fact"])[:400]} for item in items
            ]
            groups = _batches(reduction_input, max_input_bytes)
            if len(groups) >= len(items):
                raise ValueError("Evidence could not fit within the model input budget")
            reduced = await asyncio.gather(*(map_batch(group, True) for group in groups))
            items = [item.model_dump() for digest in reduced for item in digest.items]
        else:
            raise ValueError("Evidence reduction exceeded the input budget")
        evidence: object = items
    else:
        evidence = excerpts
    document = await _validated_call(
        llm,
        artifact_messages(recipe, evidence, options),
        RECIPES_REGISTRY[recipe].output_schema,
        lambda value: _validate_output(value, recipe, set(labeled), options),
    )
    used: set[int] = set()
    if isinstance(document, ArtifactDocument):
        used.update(ref for section in document.sections for ref in section.citation_refs)
    elif isinstance(document, MindMapArtifact):

        def collect(node: MindMapNode) -> None:
            used.update(node.source_refs)
            for child in node.children:
                collect(child)

        collect(document.root)
    elif isinstance(document, FlashcardsArtifact):
        used.update(card.source_ref for card in document.cards)
    elif isinstance(document, QuizArtifact):
        used.update(question.source_ref for question in document.questions)
    used_refs = sorted(used)
    citations: list[Citation] = [
        {
            "n": number,
            "chunk_id": labeled[number].id,
            "source_id": labeled[number].source_id,
            "page_start": labeled[number].page_start,
            "time_start": labeled[number].time_start,
            "time_end": labeled[number].time_end,
            "char_start": labeled[number].char_start,
            "char_end": labeled[number].char_end,
            "source_name": source_names.get(labeled[number].source_id, "Source"),
            "quote": labeled[number].text[:500],
        }
        for number in used_refs
    ]
    return cast(GeneratedArtifact, document), citations


class FakeArtifactLLM:
    """Deterministic local artifact model for tests and offline development."""

    async def stream(self, messages: list[dict[str, str]]) -> AsyncIterator[str]:
        data = json.loads(messages[1]["content"])
        if data["task"] in {"evidence_map", "evidence_reduce"}:
            yield EvidenceDigest(
                items=[
                    EvidenceItem(ref=item["ref"], fact=str(item["text"])[:200])
                    for item in data["excerpts"][:3]
                ]
            ).model_dump_json()
            return
        evidence = data["evidence"]
        first = evidence[0]
        ref = first["ref"]
        fact = str(first.get("fact") or first.get("text") or "Source finding").splitlines()[0][:250]
        artifact_cls = RECIPES_REGISTRY[data["recipe"]].output_schema
        if artifact_cls is MindMapArtifact:
            yield MindMapArtifact(
                title="Mind map",
                root=MindMapNode(
                    label="Key finding", summary=f"{fact} [S{ref}]", source_refs=[ref], children=[]
                ),
            ).model_dump_json()
            return
        if artifact_cls is FlashcardsArtifact:
            yield FlashcardsArtifact(
                title="Flashcards",
                cards=[
                    Flashcard(
                        front="What is the key finding?", back=f"{fact} [S{ref}]", source_ref=ref
                    )
                ],
            ).model_dump_json()
            return
        if artifact_cls is QuizArtifact:
            yield QuizArtifact(
                title="Quiz",
                questions=[
                    QuizQuestion(
                        question=f"Which statement is supported by the source? ({index + 1})",
                        options=[fact, "Not stated", "Opposite of the source", "None of these"],
                        answer_index=0,
                        explanation=f"The source supports this. [S{ref}]",
                        source_ref=ref,
                    )
                    for index in range(int(data.get("question_count") or 10))
                ],
            ).model_dump_json()
            return
        yield ArtifactDocument(
            recipe=data["recipe"],
            title=data["title"],
            sections=[
                ArtifactSection(
                    heading="Key finding", markdown=f"{fact} [S{ref}]", citation_refs=[ref]
                )
            ],
        ).model_dump_json()


def markdown_export(document: ArtifactDocument, citations: list[Citation]) -> str:
    lines = [f"# {document.title}", ""]
    for section in document.sections:
        lines.extend([f"## {section.heading}", "", section.markdown, ""])
    lines.extend(["## Sources", ""])
    for citation in citations:
        lines.append(
            f"[S{citation['n']}] {citation['source_name']}"
            + (f", page {citation['page_start']}" if citation["page_start"] else "")
        )
    return "\n".join(lines) + "\n"


def docx_export(document: ArtifactDocument, citations: list[Citation]) -> bytes:
    output = Document()
    output.add_heading(document.title, 0)
    for section in document.sections:
        output.add_heading(section.heading, level=1)
        output.add_paragraph(section.markdown)
    output.add_heading("Sources", level=1)
    for citation in citations:
        output.add_paragraph(
            f"[S{citation['n']}] {citation['source_name']}"
            + (f", page {citation['page_start']}" if citation["page_start"] else "")
        )
    buffer = BytesIO()
    output.save(buffer)
    return buffer.getvalue()


def pdf_export(document: ArtifactDocument, citations: list[Citation]) -> bytes:
    """Render a safe printable PDF through WeasyPrint in the Linux API image."""
    from weasyprint import HTML  # type: ignore[import-untyped]

    def render_markdown(markdown: str) -> str:
        blocks: list[str] = []
        for paragraph in markdown.split("\n\n"):
            lines = paragraph.splitlines()
            heading = re.match(r"^(#{1,6})\s+(.+)$", lines[0]) if lines else None
            if heading:
                level = len(heading.group(1))
                blocks.append(f"<h{level}>{html.escape(heading.group(2))}</h{level}>")
            elif lines and all(re.match(r"^\s*[-*]\s+", line) for line in lines):
                items = []
                for line in lines:
                    item = re.sub(r"^\s*[-*]\s+", "", line)
                    items.append(f"<li>{html.escape(item)}</li>")
                blocks.append(f"<ul>{''.join(items)}</ul>")
            elif paragraph.strip():
                escaped = html.escape(paragraph).replace("\n", "<br>")
                escaped = re.sub(r"\*\*(.+?)\*\*", r"<strong>\1</strong>", escaped)
                escaped = re.sub(r"`([^`]+)`", r"<code>\1</code>", escaped)
                blocks.append(f"<p>{escaped}</p>")
        return "\n".join(blocks)

    sections = "\n".join(
        "<section><h2>"
        + html.escape(section.heading)
        + "</h2>"
        + render_markdown(section.markdown)
        + "</section>"
        for section in document.sections
    )
    sources = []
    for citation in citations:
        page = f", page {citation['page_start']}" if citation["page_start"] else ""
        sources.append(
            f"<li>[S{citation['n']}] {html.escape(citation['source_name'])}{html.escape(page)}</li>"
        )
    markup = f"""<!doctype html><html><head><meta charset="utf-8"><style>
        @page {{ size: A4; margin: 18mm 16mm; }}
        body {{ color: #172033; font: 11pt/1.55 sans-serif; }}
        h1 {{ font-size: 24pt; }} h2 {{ margin-top: 1.8em; border-bottom: 1px solid #ccd; }}
        section {{ break-inside: avoid; }} code {{ font-family: monospace; }}
        </style></head><body><h1>{html.escape(document.title)}</h1>{sections}
        <h2>Sources</h2><ul>{"".join(sources)}</ul></body></html>"""
    return cast(bytes, HTML(string=markup).write_pdf())
